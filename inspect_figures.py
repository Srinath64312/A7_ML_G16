import docx, sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('new srinath 106 THINK-AI.docx')
for i, p in enumerate(doc.paragraphs):
    drawings = p._element.xpath('.//w:drawing')
    if drawings:
        next_text = doc.paragraphs[i+1].text if i+1 < len(doc.paragraphs) else "END"
        print(f"P{i}: Drawing found! -> Caption at P{i+1}: '{next_text}'")
